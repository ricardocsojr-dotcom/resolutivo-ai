import os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ===== TOKENS (mesma especificação visual do HTML) =====
C_TEXT      = RGBColor(0x63, 0x66, 0x6A)   # cinza tipográfico oficial
C_BLACK     = RGBColor(0x00, 0x00, 0x00)
C_ACCENT    = RGBColor(0xF7, 0xA8, 0x00)   # laranja RD
C_ACCENT_LT = RGBColor(0xFF, 0xF8, 0xEB)
C_CARD_BG   = RGBColor(0xF8, 0xFA, 0xFC)
C_BORDER    = RGBColor(0xE2, 0xE8, 0xF0)
C_GREEN     = RGBColor(0x04, 0x78, 0x57)
C_GREEN_BG  = RGBColor(0xEC, 0xFD, 0xF5)
C_RED       = RGBColor(0xB9, 0x1C, 0x1C)
C_RED_BG    = RGBColor(0xFE, 0xF2, 0xF2)
C_NEUTRAL_BG= RGBColor(0xEE, 0xF1, 0xF4)
C_WHITE     = RGBColor(0xFF, 0xFF, 0xFF)

FONT = "Lato"
LOGO_PATH = "C:/Projetos/resolutivo-ai/skills/romano-donadel-slide-style/assets/logo_romano_donadel.png"
OUTPUT_PPTX = "C:/Projetos/resolutivo-ai/rescisoria_trivale_minare_2slides.pptx"

prs = Presentation()
prs.slide_width = Emu(12192000)   # 16:9
prs.slide_height = Emu(6858000)
BLANK = prs.slide_layouts[6]

def tb(slide, x, y, w, h):
    box = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    return tf

def para(tf, text, size, bold=False, color=C_TEXT, align=PP_ALIGN.LEFT, space_before=0, space_after=0, first=False):
    p = tf.paragraphs[0] if (first and not tf.paragraphs[0].runs) else tf.add_paragraph()
    p.text = text
    p.font.name = FONT
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = color
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    return p

def rect(slide, x, y, w, h, fill, line=None, line_w=None):
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    shp.fill.solid()
    shp.fill.fore_color.rgb = fill
    if line is not None:
        shp.line.color.rgb = line
        shp.line.width = Pt(line_w or 1)
    else:
        shp.line.fill.background()
    shp.shadow.inherit = False
    return shp

def add_header(slide, eyebrow, title, page_num):
    # accent bar
    rect(slide, 594360, 380000, 55000, 460000, C_ACCENT)
    tf = tb(slide, 730000, 320000, 9600000, 620000)
    para(tf, eyebrow.upper(), 10.5, bold=True, color=C_ACCENT, first=True)
    para(tf, title.upper(), 19, bold=True, color=C_BLACK, space_before=2)
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Emu(10480000), Emu(330000), width=Emu(1180000))
    # footer divider + text
    rect(slide, 594360, 6440000, 11000000, 8000, C_BORDER)
    tf_f = tb(slide, 594360, 6480000, 6000000, 240000)
    para(tf_f, "CONFIDENCIAL | ROMANO DONADEL ADVOGADOS ASSOCIADOS", 8.5, bold=True, color=C_TEXT, first=True)
    tf_p = tb(slide, 10500000, 6480000, 1100000, 240000)
    para(tf_p, f"SLIDE {page_num:02d} / 02", 8.5, bold=True, color=C_BLACK, align=PP_ALIGN.RIGHT, first=True)

def add_core_idea(slide, x, y, w, headline, support):
    """Ideia central como bloco autônomo — escala grande, isolado, barra lateral laranja."""
    rect(slide, x, y, 55000, 620000, C_ACCENT)
    tf = tb(slide, x + 180000, y - 20000, w - 200000, 660000)
    para(tf, headline, 19, bold=True, color=C_BLACK, first=True)
    para(tf, support, 11, bold=False, color=C_TEXT, space_before=6)

# ==========================================================
# SLIDE 1 — MECANISMO
# ==========================================================
s1 = prs.slides.add_slide(BLANK)
add_header(s1, "Mecanismo — Ação Rescisória nº 1.0000.26.447718-3/000",
           "Cada Repasse Deixa um Rastro Documental Completo", 1)

# Bloco 1: fluxo operacional
tf_h1 = tb(s1, 594360, 1150000, 6000000, 300000)
para(tf_h1, "1  FLUXO OPERACIONAL — O DINHEIRO", 12, bold=True, color=C_BLACK, first=True)

flow = [
    ("MINARE (LOJA)", "Venda no cartão", "Cliente final paga na Minare com o cartão da rede Valecard."),
    ("TRIVALE (REDE)", "Apuração", "Trivale registra a transação e desconta a taxa (MDR)."),
    ("TRIVALE → MINARE", "Repasse", "Valor líquido transferido por TED/DOC, individual ou agrupado."),
    ("TRIVALE (DOCUMENTA)", "Registro", "Extrato de Conveniada lista quais títulos aquele repasse quitou."),
]
card_w = 2620000
gap = 100000
x0 = 594360
y0 = 1500000
card_h = 1350000
for i, (actor, title, desc) in enumerate(flow):
    x = x0 + i * (card_w + gap)
    rect(s1, x, y0, card_w, card_h, C_CARD_BG, line=C_BORDER, line_w=1.2)
    tf = tb(s1, x + 150000, y0 + 130000, card_w - 300000, card_h - 260000)
    # actor pill (simplificado como texto laranja pequeno)
    para(tf, actor, 8.5, bold=True, color=RGBColor(0xB4, 0x53, 0x09), first=True)
    para(tf, title.upper(), 12.5, bold=True, color=C_BLACK, space_before=4)
    para(tf, desc, 9.5, bold=False, color=C_TEXT, space_before=4)
    if i < len(flow) - 1:
        arrow = s1.shapes.add_textbox(Emu(x + card_w + 5000), Emu(y0 + card_h/2 - 100000), Emu(gap + 20000), Emu(200000))
        atf = arrow.text_frame
        ap = atf.paragraphs[0]
        ap.text = "➔"
        ap.font.size = Pt(16)
        ap.font.bold = True
        ap.font.color.rgb = C_ACCENT
        ap.alignment = PP_ALIGN.CENTER

# Bloco 2: cadeia documental
tf_h2 = tb(s1, 594360, 3050000, 6000000, 300000)
para(tf_h2, "2  CADEIA DOCUMENTAL — A PROVA, TÍTULO A TÍTULO", 12, bold=True, color=C_BLACK, first=True)

docs = [
    ("ETAPA 1", "Título de venda", "Prova que a obrigação existiu — o que o laudo listou como \"em aberto\"."),
    ("ETAPA 2", "Extrato de Conveniada", "Prova que aquele título foi liquidado — valor bruto e líquido."),
    ("ETAPA 3", "Agrupamento bancário", "Quando aplicável: vários títulos pagos num único lote."),
    ("ETAPA 4", "TED / DOC", "Prova bancária final: o valor efetivamente saiu e chegou."),
]
y1 = 3400000
card_h2 = 1250000
for i, (tag, title, desc) in enumerate(docs):
    x = x0 + i * (card_w + gap)
    rect(s1, x, y1, card_w, card_h2, C_WHITE, line=C_BLACK, line_w=1.4)
    tf = tb(s1, x + 150000, y1 + 130000, card_w - 300000, card_h2 - 260000)
    para(tf, tag, 8.5, bold=True, color=C_ACCENT, first=True)
    para(tf, title.upper(), 12, bold=True, color=C_BLACK, space_before=4)
    para(tf, desc, 9.5, bold=False, color=C_TEXT, space_before=4)
    if i < len(docs) - 1:
        arrow = s1.shapes.add_textbox(Emu(x + card_w + 5000), Emu(y1 + card_h2/2 - 100000), Emu(gap + 20000), Emu(200000))
        atf = arrow.text_frame
        ap = atf.paragraphs[0]
        ap.text = "⟶"
        ap.font.size = Pt(14)
        ap.font.color.rgb = C_BLACK
        ap.alignment = PP_ALIGN.CENTER

# Ideia central (destaque real, não legenda)
add_core_idea(
    s1, 594360, 5150000, 11000000,
    "A quitação só está provada quando as 4 etapas se conectam.",
    "Título, extrato, agrupamento (quando houver) e comprovante bancário — nenhuma etapa isolada basta."
)

# ==========================================================
# SLIDE 2 — ERRO DE FATO (decomposição parte-todo)
# ==========================================================
s2 = prs.slides.add_slide(BLANK)
add_header(s2, "Erro de Fato — CPC, art. 966, VIII e §1º",
           "Dentro da Dívida Apontada, Quanto Já Foi Pago?", 2)

# Âncora: valor total tratado como devido
tf_anchor = tb(s2, 594360, 1150000, 8000000, 400000)
p1 = tf_anchor.paragraphs[0]
p1.text = "Valor que a sentença tratou como devido:  "
p1.font.name = FONT
p1.font.size = Pt(12.5)
p1.font.bold = True
p1.font.color.rgb = C_TEXT
run2 = p1.add_run()
run2.text = "R$ 150.730,93"
run2.font.name = FONT
run2.font.size = Pt(24)
run2.font.bold = True
run2.font.color.rgb = C_BLACK

# Barra única de decomposição (parte-todo) — proporcional aos 3 segmentos
bar_x = 594360
bar_y = 1650000
bar_w = 11000000
bar_h = 700000
pago_w = int(bar_w * 0.852)
dup_w = int(bar_w * 0.025)
aberto_w = bar_w - pago_w - dup_w

# outer border frame
outer = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(bar_x), Emu(bar_y), Emu(bar_w), Emu(bar_h))
outer.fill.background()
outer.line.color.rgb = C_BLACK
outer.line.width = Pt(2)
outer.shadow.inherit = False

seg_pago = rect(s2, bar_x, bar_y, pago_w, bar_h, C_GREEN_BG)
seg_dup = rect(s2, bar_x + pago_w, bar_y, dup_w, bar_h, C_NEUTRAL_BG)
seg_aberto = rect(s2, bar_x + pago_w + dup_w, bar_y, aberto_w, bar_h, C_RED_BG)

tf_seg1 = tb(s2, bar_x + 100000, bar_y + bar_h/2 - 140000, pago_w - 200000, 280000)
para(tf_seg1, "✓ JÁ PAGO — R$ 128.370,50", 13, bold=True, color=C_GREEN, align=PP_ALIGN.CENTER, first=True)

tf_seg3 = tb(s2, bar_x + pago_w + dup_w + 50000, bar_y + bar_h/2 - 140000, aberto_w - 100000, 280000)
para(tf_seg3, "R$ 18.523,21", 13, bold=True, color=C_RED, align=PP_ALIGN.CENTER, first=True)

tf_note = tb(s2, bar_x, bar_y + bar_h + 60000, bar_w, 300000)
para(tf_note, "Faixa cinza (R$ 3.837,22): valor contado duas vezes no laudo — não é dívida nem pagamento, é erro de soma já descontado nos números ao lado.",
     9.5, bold=False, color=RGBColor(0x94, 0x9E, 0xB0), first=True)

# Callouts (dois cartões)
call_y = 2900000
call_w = 5350000
call_h = 1550000

rect(s2, 594360, call_y, call_w, call_h, C_GREEN_BG)
tf_c1 = tb(s2, 594360 + 200000, call_y + 180000, call_w - 400000, call_h - 360000)
para(tf_c1, "DENTRO DO VALOR DEVIDO: JÁ PAGO", 10.5, bold=True, color=C_GREEN, first=True)
para(tf_c1, "R$ 128.370,50", 19, bold=True, color=C_BLACK, space_before=4)
para(tf_c1, "85% do que a sentença tratou como dívida já tinha comprovante nos autos — título, extrato e TED/DOC — antes mesmo da sentença.", 10, bold=False, color=C_TEXT, space_before=4)

rect(s2, 6244360, call_y, call_w, call_h, C_RED_BG)
tf_c2 = tb(s2, 6244360 + 200000, call_y + 180000, call_w - 400000, call_h - 360000)
para(tf_c2, "DENTRO DO VALOR DEVIDO: REALMENTE EM ABERTO", 10.5, bold=True, color=C_RED, first=True)
para(tf_c2, "R$ 18.523,21", 19, bold=True, color=C_BLACK, space_before=4)
para(tf_c2, "Única fração que sobra como dívida real depois de descontar o que já foi pago. A Trivale não contesta esse valor.", 10, bold=False, color=C_TEXT, space_before=4)

# Ideia central (destaque real)
add_core_idea(
    s2, 594360, 4900000, 11000000,
    "85% do valor tratado como dívida já estava pago e documentado.",
    "A sentença nunca cruzou o título com o comprovante — isso é erro de fato (CPC, art. 966, VIII), não rediscussão de mérito."
)

prs.save(OUTPUT_PPTX)
print(f"PPTX gerado a partir da mesma especificação visual do HTML: {OUTPUT_PPTX}")
