import os
from pptx import Presentation
from pptx.util import Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# Core Brand Tokens
COLOR_BG = RGBColor(255, 255, 255)
COLOR_DARK_SLATE = RGBColor(15, 23, 42)        # #0F172A
COLOR_TEXT_PRIMARY = RGBColor(15, 23, 42)
COLOR_STRUCTURE = RGBColor(99, 102, 106)       # #63666A
COLOR_ACCENT = RGBColor(247, 168, 0)            # #F7A800
COLOR_CARD_BG = RGBColor(248, 250, 252)
COLOR_CARD_BORDER = RGBColor(226, 232, 240)
COLOR_ACCENT_LIGHT = RGBColor(255, 248, 235)
COLOR_GREEN_BG = RGBColor(236, 253, 245)
COLOR_GREEN_TEXT = RGBColor(4, 120, 87)
COLOR_RED_BG = RGBColor(254, 242, 242)
COLOR_RED_TEXT = RGBColor(185, 28, 28)

FONT_LATO = "Lato"
LOGO_PATH = "C:/Projetos/resolutivo-ai/skills/romano-donadel-slide-style/assets/logo_romano_donadel.png"
OUTPUT_PPTX = "C:/Projetos/resolutivo-ai/Apresentacao_Caso_Silvio_Afonso_Romano_Donadel.pptx"

prs = Presentation()
prs.slide_width = Emu(12192000)   # 16:9 widescreen
prs.slide_height = Emu(6858000)
blank_layout = prs.slide_layouts[6]

def add_header(slide, title_text, category_text=""):
    rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(400000), Emu(60000), Emu(450000))
    rule.fill.solid()
    rule.fill.fore_color.rgb = COLOR_ACCENT
    rule.line.color.rgb = COLOR_ACCENT
    
    tb = slide.shapes.add_textbox(Emu(740000), Emu(340000), Emu(9500000), Emu(600000))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    
    if category_text:
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = FONT_LATO
        p_cat.font.size = Pt(9.5)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT
        p_title = tf.add_paragraph()
    else:
        p_title = tf.paragraphs[0]
        
    p_title.text = title_text.upper()
    p_title.font.name = FONT_LATO
    p_title.font.size = Pt(17)
    p_title.font.bold = True
    p_title.font.color.rgb = COLOR_DARK_SLATE
    
    if os.path.exists(LOGO_PATH):
        slide.shapes.add_picture(LOGO_PATH, Emu(10500000), Emu(320000), width=Emu(1150000))

def add_footer(slide, page_num, total_pages=7):
    div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(6420000), Emu(11000000), Emu(8000))
    div.fill.solid()
    div.fill.fore_color.rgb = COLOR_CARD_BORDER
    div.line.color.rgb = COLOR_CARD_BORDER
    
    tb_l = slide.shapes.add_textbox(Emu(594360), Emu(6460000), Emu(5500000), Emu(250000))
    tf_l = tb_l.text_frame
    tf_l.margin_left = tf_l.margin_top = tf_l.margin_right = tf_l.margin_bottom = 0
    p_l = tf_l.paragraphs[0]
    p_l.text = "CONFIDENCIAL | ROMANO DONADEL ADVOGADOS ASSOCIADOS"
    p_l.font.name = FONT_LATO
    p_l.font.size = Pt(8.5)
    p_l.font.color.rgb = COLOR_STRUCTURE
    
    tb_r = slide.shapes.add_textbox(Emu(10500000), Emu(6460000), Emu(1100000), Emu(250000))
    tf_r = tb_r.text_frame
    tf_r.margin_left = tf_r.margin_top = tf_r.margin_right = tf_r.margin_bottom = 0
    p_r = tf_r.paragraphs[0]
    p_r.text = f"SLIDE {page_num:02d} / {total_pages:02d}"
    p_r.alignment = PP_ALIGN.RIGHT
    p_r.font.name = FONT_LATO
    p_r.font.size = Pt(8.5)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_DARK_SLATE

# ==========================================
# SLIDE 1: CAPA EDITORIAL
# ==========================================
s1 = prs.slides.add_slide(blank_layout)

# Left Column (Dark Slate)
left_panel = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(600000), Emu(4500000), Emu(5600000))
left_panel.fill.solid()
left_panel.fill.fore_color.rgb = COLOR_DARK_SLATE
left_panel.line.color.rgb = COLOR_DARK_SLATE

# Orange Accent Strip at bottom of left panel
accent_strip = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(6140000), Emu(4500000), Emu(60000))
accent_strip.fill.solid()
accent_strip.fill.fore_color.rgb = COLOR_ACCENT
accent_strip.line.color.rgb = COLOR_ACCENT

# Text inside Left Panel
tb_lp = s1.shapes.add_textbox(Emu(850000), Emu(900000), Emu(3950000), Emu(5000000))
tf_lp = tb_lp.text_frame
tf_lp.word_wrap = True

p = tf_lp.paragraphs[0]
p.text = "ROMANO DONADEL ADVOGADOS"
p.font.name = FONT_LATO
p.font.size = Pt(11)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

p_sub = tf_lp.add_paragraph()
p_sub.text = "Defesa Patrimonial & Blindagem de Meação"
p_sub.font.name = FONT_LATO
p_sub.font.size = Pt(16)
p_sub.font.bold = True
p_sub.font.color.rgb = RGBColor(255, 255, 255)
p_sub.space_before = Pt(6)
p_sub.space_after = Pt(28)

kpis = [
    ("03 FRENTES", "Ações Judiciais Coordenadas"),
    ("R$ 0 DE MULTA", "Astreintes Revogadas no TJMG"),
    ("100% BLINDADO", "Apartamento Sense Preservado")
]
for val, lbl in kpis:
    pk = tf_lp.add_paragraph()
    pk.text = val
    pk.font.name = FONT_LATO
    pk.font.size = Pt(20)
    pk.font.bold = True
    pk.font.color.rgb = COLOR_ACCENT
    pk.space_before = Pt(12)
    
    pl = tf_lp.add_paragraph()
    pl.text = lbl
    pl.font.name = FONT_LATO
    pl.font.size = Pt(9.5)
    pl.font.color.rgb = RGBColor(203, 213, 225)

# Right Column
if os.path.exists(LOGO_PATH):
    s1.shapes.add_picture(LOGO_PATH, Emu(10300000), Emu(700000), width=Emu(1400000))

tb_rp = s1.shapes.add_textbox(Emu(5500000), Emu(1800000), Emu(6100000), Emu(3800000))
tf_rp = tb_rp.text_frame
tf_rp.word_wrap = True

pr1 = tf_rp.paragraphs[0]
pr1.text = "RELATÓRIO EXECUTIVO RDAA"
pr1.font.name = FONT_LATO
pr1.font.size = Pt(10)
pr1.font.bold = True
pr1.font.color.rgb = COLOR_STRUCTURE

pr2 = tf_rp.add_paragraph()
pr2.text = "CASO SÍLVIO LUIZ AFONSO"
pr2.font.name = FONT_LATO
pr2.font.size = Pt(32)
pr2.font.bold = True
pr2.font.color.rgb = COLOR_DARK_SLATE
pr2.space_before = Pt(8)

pr3 = tf_rp.add_paragraph()
pr3.text = "Estratégia Processual Integrada, Responsabilidade Tributária do Imóvel Ocupado e Governança de Acervo"
pr3.font.name = FONT_LATO
pr3.font.size = Pt(14)
pr3.font.color.rgb = COLOR_STRUCTURE
pr3.space_before = Pt(6)
pr3.space_after = Pt(28)

pr4 = tf_rp.add_paragraph()
pr4.text = "👤 CLIENTE: Sílvio Luiz Afonso   |   ⚖ PARTE ADVERSA: Ana Lúcia de Oliveira Afonso"
pr4.font.name = FONT_LATO
pr4.font.size = Pt(11)
pr4.font.bold = True
pr4.font.color.rgb = COLOR_DARK_SLATE

pr5 = tf_rp.add_paragraph()
pr5.text = "PATROCÍNIO: Romano Donadel Advogados Associados   |   DATA: Setembro / 2026"
pr5.font.name = FONT_LATO
pr5.font.size = Pt(10)
pr5.font.color.rgb = COLOR_STRUCTURE
pr5.space_before = Pt(4)

add_footer(s1, 1, 7)

# ==========================================
# SLIDE 2: BENTO GRID EXECUTIVO
# ==========================================
s2 = prs.slides.add_slide(blank_layout)
add_header(s2, "Arquitetura Estratégica do Contencioso", "Visão Geral & Governança")
add_footer(s2, 2, 7)

# Left Large Dark Card
b_left = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(1200000), Emu(3800000), Emu(4900000))
b_left.fill.solid()
b_left.fill.fore_color.rgb = COLOR_DARK_SLATE
b_left.line.color.rgb = COLOR_DARK_SLATE

tb_bl = s2.shapes.add_textbox(Emu(800000), Emu(1400000), Emu(3400000), Emu(4500000))
tf_bl = tb_bl.text_frame
tf_bl.word_wrap = True

p = tf_bl.paragraphs[0]
p.text = "GESTÃO COORDENADA"
p.font.name = FONT_LATO
p.font.size = Pt(9.5)
p.font.bold = True
p.font.color.rgb = COLOR_ACCENT

p_num = tf_bl.add_paragraph()
p_num.text = "03"
p_num.font.name = FONT_LATO
p_num.font.size = Pt(44)
p_num.font.bold = True
p_num.font.color.rgb = COLOR_ACCENT

p_t = tf_bl.add_paragraph()
p_t.text = "FRENTES PROCESSUAIS"
p_t.font.name = FONT_LATO
p_t.font.size = Pt(13)
p_t.font.bold = True
p_t.font.color.rgb = RGBColor(255, 255, 255)
p_t.space_after = Pt(14)

bullets_left = [
    "3ª Vara de Família: Divórcio & Partilha de Haveres",
    "TJMG (8ª Câmara): Agravo de IPTU Begônias",
    "10ª Vara Cível: Arbitramento de Aluguel"
]
for b in bullets_left:
    pb = tf_bl.add_paragraph()
    pb.text = f"• {b}"
    pb.font.name = FONT_LATO
    pb.font.size = Pt(10)
    pb.font.color.rgb = RGBColor(241, 245, 249)
    pb.space_after = Pt(6)

# 4 Right Cards (2x2 Grid)
grid_cards = [
    ("ACERVOS", "02 IMÓVEIS CENTRAIS", "• Rua das Begônias: Fruição exclusiva de Ana\n• Ed. Sense Vertical: Aquisição autônoma de Sílvio", COLOR_CARD_BG, COLOR_CARD_BORDER),
    ("DECISÃO TJMG", "VITÓRIA RECURSAL AI /006", "• Astreintes de R$ 10k revogadas\n• IPTU Begônias: Obrigação integral mantida para Ana", COLOR_ACCENT_LIGHT, COLOR_ACCENT),
    ("10ª VARA CÍVEL", "SANEADOR ART. 357, §1º", "• Fato Gerador: Dispensa de esbulho físico\n• Termo Inicial: Piso garantido na citação", COLOR_CARD_BG, COLOR_CARD_BORDER),
    ("ALIMENTOS", "SEGREGAÇÃO ESTRITA", "• Bloqueio a qualquer compensação cruzada\n• Execução de 4,5 SM mantida em via autônoma", COLOR_CARD_BG, COLOR_CARD_BORDER)
]

gw = Emu(3450000)
gh = Emu(2350000)
coords = [
    (Emu(4650000), Emu(1200000)),
    (Emu(8300000), Emu(1200000)),
    (Emu(4650000), Emu(3750000)),
    (Emu(8300000), Emu(3750000))
]

for idx, (tag, title, body, bg_col, bdr_col) in enumerate(grid_cards):
    gx, gy = coords[idx]
    box = s2.shapes.add_shape(MSO_SHAPE.RECTANGLE, gx, gy, gw, gh)
    box.fill.solid()
    box.fill.fore_color.rgb = bg_col
    box.line.color.rgb = bdr_col
    
    tbg = s2.shapes.add_textbox(gx + Emu(180000), gy + Emu(160000), gw - Emu(360000), gh - Emu(320000))
    tfg = tbg.text_frame
    tfg.word_wrap = True
    
    pt = tfg.paragraphs[0]
    pt.text = tag
    pt.font.name = FONT_LATO
    pt.font.size = Pt(8.5)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_ACCENT
    
    ptt = tfg.add_paragraph()
    ptt.text = title
    ptt.font.name = FONT_LATO
    ptt.font.size = Pt(12)
    ptt.font.bold = True
    ptt.font.color.rgb = COLOR_DARK_SLATE
    ptt.space_after = Pt(8)
    
    for line in body.split("\n"):
        pbl = tfg.add_paragraph()
        pbl.text = line
        pbl.font.name = FONT_LATO
        pbl.font.size = Pt(9.5)
        pbl.font.color.rgb = COLOR_DARK_SLATE
        pbl.space_after = Pt(3)

# ==========================================
# SLIDE 3: MATRIZ DE DEMANDAS CRUZADAS
# ==========================================
s3 = prs.slides.add_slide(blank_layout)
add_header(s3, "Mapeamento Dinâmico das Demandas Cruzadas", "Estrutura de Litigância")
add_footer(s3, 3, 7)

# 4 Horizontal Process Pillars
proc_pillars = [
    ("VARA DE FAMÍLIA", "Divórcio e Partilha", "Proc. 5026681-10", "Partilha do patrimônio comum e apuração de haveres.", "Isolar bens particulares e exigir prestação de contas dos frutos."),
    ("TJMG — 8ª CÂMARA", "AI IPTU Begônias (/006)", "Proc. 1.0000.23.223786-7", "Custeio do IPTU e baixa de protesto da casa residencial.", "Manter dever tributário exclusivo da coproprietária ocupante."),
    ("10ª VARA CÍVEL", "Arbitramento de Aluguéis", "Proc. 5033450-63", "Indenização pela fruição exclusiva da residência.", "Fixar piso incontroverso na citação e dispensar esbulho físico."),
    ("10ª VARA CÍVEL", "Reconvenção (Ed. Sense)", "Proc. 5033450-63", "Pretensão de meação sobre o apartamento de Sílvio.", "Impor ônus primário integral à reconvinte (art. 373, I do CPC).")
]

pw3 = Emu(2600000)
for idx, (court, action, proc, desc, target) in enumerate(proc_pillars):
    px = Emu(594360 + idx * 2800000)
    p_box = s3.shapes.add_shape(MSO_SHAPE.RECTANGLE, px, Emu(1200000), pw3, Emu(4900000))
    p_box.fill.solid()
    p_box.fill.fore_color.rgb = COLOR_CARD_BG
    p_box.line.color.rgb = COLOR_CARD_BORDER
    
    tbp = s3.shapes.add_textbox(px + Emu(160000), Emu(1380000), pw3 - Emu(320000), Emu(4500000))
    tfp = tbp.text_frame
    tfp.word_wrap = True
    
    pc = tfp.paragraphs[0]
    pc.text = court
    pc.font.name = FONT_LATO
    pc.font.size = Pt(8.5)
    pc.font.bold = True
    pc.font.color.rgb = COLOR_ACCENT
    
    pa = tfp.add_paragraph()
    pa.text = action
    pa.font.name = FONT_LATO
    pa.font.size = Pt(12)
    pa.font.bold = True
    pa.font.color.rgb = COLOR_DARK_SLATE
    
    pp = tfp.add_paragraph()
    pp.text = proc
    pp.font.name = FONT_LATO
    pp.font.size = Pt(8.5)
    pp.font.color.rgb = COLOR_STRUCTURE
    pp.space_after = Pt(12)
    
    pd = tfp.add_paragraph()
    pd.text = desc
    pd.font.name = FONT_LATO
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = COLOR_DARK_SLATE
    pd.space_after = Pt(16)
    
    pt = tfp.add_paragraph()
    pt.text = f"🎯 DIRETRIZ RDAA:\n{target}"
    pt.font.name = FONT_LATO
    pt.font.size = Pt(9.5)
    pt.font.bold = True
    pt.font.color.rgb = COLOR_DARK_SLATE

# ==========================================
# SLIDE 4: FRENTE 1 IPTU BEGÔNIAS
# ==========================================
s4 = prs.slides.add_slide(blank_layout)
add_header(s4, "TJMG Mantém IPTU com Ocupante e Revoga Multa de R$ 10k", "AI 1.0000.23.223786-7/006 — TJMG")
add_footer(s4, 4, 7)

# Flow Strip
steps_flow = [
    ("PASSO 1: Fruição Exclusiva", "Ana reside no imóvel desde 13/03/2023"),
    ("PASSO 2: Inadimplência", "IPTU e taxas deixados em aberto"),
    ("PASSO 3: Protesto Indevido", "Fazenda protestou Sílvio em cartório"),
    ("PASSO 4: Ordem TJMG", "Ana deve solver dívida e baixar protesto")
]
fw = Emu(2600000)
for idx, (h, b) in enumerate(steps_flow):
    fx = Emu(594360 + idx * 2800000)
    fb = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, fx, Emu(1200000), fw, Emu(1300000))
    fb.fill.solid()
    if idx == 3:
        fb.fill.fore_color.rgb = COLOR_GREEN_BG
        fb.line.color.rgb = COLOR_GREEN_TEXT
    else:
        fb.fill.fore_color.rgb = COLOR_CARD_BG
        fb.line.color.rgb = COLOR_CARD_BORDER
        
    tbf = s4.shapes.add_textbox(fx + Emu(120000), Emu(1300000), fw - Emu(240000), Emu(1100000))
    tff = tbf.text_frame
    tff.word_wrap = True
    
    pf1 = tff.paragraphs[0]
    pf1.text = h
    pf1.font.name = FONT_LATO
    pf1.font.size = Pt(9.5)
    pf1.font.bold = True
    pf1.font.color.rgb = COLOR_GREEN_TEXT if idx == 3 else COLOR_ACCENT
    
    pf2 = tff.add_paragraph()
    pf2.text = b
    pf2.font.name = FONT_LATO
    pf2.font.size = Pt(9)
    pf2.font.color.rgb = COLOR_DARK_SLATE

# 2 Contrast Boxes below
cw = Emu(5350000)
ch = Emu(3300000)
cy = Emu(2800000)

# Adversary (Danger)
box_adv = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), cy, cw, ch)
box_adv.fill.solid()
box_adv.fill.fore_color.rgb = COLOR_RED_BG
box_adv.line.color.rgb = COLOR_RED_TEXT

tb_adv = s4.shapes.add_textbox(Emu(850000), cy + Emu(220000), cw - Emu(500000), ch - Emu(440000))
tf_adv = tb_adv.text_frame
tf_adv.word_wrap = True

pa1 = tf_adv.paragraphs[0]
pa1.text = "✗ TESE ADVERSÁRIA REJEITADA"
pa1.font.name = FONT_LATO
pa1.font.size = Pt(11)
pa1.font.bold = True
pa1.font.color.rgb = COLOR_RED_TEXT
pa1.space_after = Pt(8)

pa2 = tf_adv.add_paragraph()
pa2.text = "• Alegação de 'violência patrimonial' e suposta dependência econômica.\n• Tentou compensar IPTU com pensão alimentícia em atraso.\n• Inconsistência: Sílvio é a vítima do protesto; alimentos têm rito autônomo."
pa2.font.name = FONT_LATO
pa2.font.size = Pt(10)
pa2.font.color.rgb = COLOR_DARK_SLATE

# Solution RDAA (Accent Light)
box_sol = s4.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(6244360), cy, cw, ch)
box_sol.fill.solid()
box_sol.fill.fore_color.rgb = COLOR_ACCENT_LIGHT
box_sol.line.color.rgb = COLOR_ACCENT
box_sol.line.width = Pt(1.5)

tb_sol = s4.shapes.add_textbox(Emu(6500000), cy + Emu(220000), cw - Emu(500000), ch - Emu(440000))
tf_sol = tb_sol.text_frame
tf_sol.word_wrap = True

ps1 = tf_sol.paragraphs[0]
ps1.text = "✔ SOLUÇÃO RDAA ACOLHIDA NO TJMG"
ps1.font.name = FONT_LATO
ps1.font.size = Pt(11)
ps1.font.bold = True
ps1.font.color.rgb = COLOR_ACCENT
ps1.space_after = Pt(8)

ps2 = tf_sol.add_paragraph()
ps2.text = "• Precedentes pacíficos: TJMG AC 0013168-49 e STJ AREsp 2.462.038.\n• O coproprietário ocupante responde pelas despesas de conservação e tributos.\n• Vitória: Mantido dever exclusivo de Ana e astreintes de R$ 10k revogadas."
ps2.font.name = FONT_LATO
ps2.font.size = Pt(10)
ps2.font.color.rgb = COLOR_DARK_SLATE

# ==========================================
# SLIDE 5: MATRIZ 2X2 SANEADOR
# ==========================================
s5 = prs.slides.add_slide(blank_layout)
add_header(s5, "Matriz de Decisão: Ajuste do Saneador (Art. 357, §1º)", "Ação Ordinária nº 5033450-63.2025.8.13.0702")
add_footer(s5, 5, 7)

m_coords = [
    (Emu(594360), Emu(1200000), "1. EXIGÊNCIA DO SANEADOR (VÍCIO)", "Decisão exigiu prova de que Ana impedia fisicamente Sílvio de ingressar no imóvel. Requisito desnecessário e prejudicial.", COLOR_RED_BG, COLOR_RED_TEXT),
    (Emu(6244360), Emu(1200000), "2. PADRÃO STJ (RESP 1.699.013/DF)", "A impossibilidade prática de coabitação associada à oposição inequívoca basta para gerar a indenização locatícia.", COLOR_ACCENT_LIGHT, COLOR_ACCENT),
    (Emu(594360), Emu(3750000), "3. PISO MÍNIMO INCONTROVERSO", "Data da citação válida como marco garantido e seguro de oposição, sem necessidade de instrução adicional.", COLOR_GREEN_BG, COLOR_GREEN_TEXT),
    (Emu(6244360), Emu(3750000), "4. RETROATIVIDADE PRETENDIDA", "Direito resguardado de demonstrar termo inicial desde 13/03/2023 através dos documentos da ação de divórcio.", COLOR_CARD_BG, COLOR_CARD_BORDER)
]

for mx, my, mtitle, mbody, mbg, mbdr in m_coords:
    mbox = s5.shapes.add_shape(MSO_SHAPE.RECTANGLE, mx, my, Emu(5350000), Emu(2350000))
    mbox.fill.solid()
    mbox.fill.fore_color.rgb = mbg
    mbox.line.color.rgb = mbdr
    
    tbm = s5.shapes.add_textbox(mx + Emu(200000), my + Emu(180000), Emu(4950000), Emu(2000000))
    tfm = tbm.text_frame
    tfm.word_wrap = True
    
    pm1 = tfm.paragraphs[0]
    pm1.text = mtitle
    pm1.font.name = FONT_LATO
    pm1.font.size = Pt(11)
    pm1.font.bold = True
    pm1.font.color.rgb = mbdr
    pm1.space_after = Pt(8)
    
    pm2 = tfm.add_paragraph()
    pm2.text = mbody
    pm2.font.name = FONT_LATO
    pm2.font.size = Pt(10)
    pm2.font.color.rgb = COLOR_DARK_SLATE

# ==========================================
# SLIDE 6: BLINDAGEM DO SENSE
# ==========================================
s6 = prs.slides.add_slide(blank_layout)
add_header(s6, "Blindagem do Apartamento Sense Vertical Living", "Defesa na Reconvenção — 10ª Vara Cível")
add_footer(s6, 6, 7)

# Adversary Attack Box
b_att = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(594360), Emu(1200000), Emu(5100000), Emu(4900000))
b_att.fill.solid()
b_att.fill.fore_color.rgb = COLOR_CARD_BG
b_att.line.color.rgb = COLOR_CARD_BORDER

tb_at = s6.shapes.add_textbox(Emu(850000), Emu(1450000), Emu(4600000), Emu(4400000))
tf_at = tb_at.text_frame
tf_at.word_wrap = True

pa_1 = tf_at.paragraphs[0]
pa_1.text = "⚠ O ATAQUE RECONVENCIONAL"
pa_1.font.name = FONT_LATO
pa_1.font.size = Pt(12)
pa_1.font.bold = True
pa_1.font.color.rgb = COLOR_STRUCTURE
pa_1.space_after = Pt(14)

pa_2 = tf_at.add_paragraph()
pa_2.text = "• Pretensão de Ana: Integrar o imóvel particular à partilha e arbitrar aluguel contra Sílvio.\n\n• Invocação do Art. 1.660 CC: Tentativa de presumir comunicabilidade sem demonstrar aporte fático.\n\n• Risco Processual: Inversão prematura do ônus da prova em detrimento do patrimônio autônomo."
pa_2.font.name = FONT_LATO
pa_2.font.size = Pt(10)
pa_2.font.color.rgb = COLOR_DARK_SLATE

# Shield Center Circle
sh_c = s6.shapes.add_shape(MSO_SHAPE.OVAL, Emu(5800000), Emu(3350000), Emu(600000), Emu(600000))
sh_c.fill.solid()
sh_c.fill.fore_color.rgb = COLOR_DARK_SLATE
sh_c.line.color.rgb = COLOR_ACCENT

# Defense RDAA Box
b_def = s6.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(6500000), Emu(1200000), Emu(5100000), Emu(4900000))
b_def.fill.solid()
b_def.fill.fore_color.rgb = COLOR_ACCENT_LIGHT
b_def.line.color.rgb = COLOR_ACCENT
b_def.line.width = Pt(1.5)

tb_df = s6.shapes.add_textbox(Emu(6750000), Emu(1450000), Emu(4600000), Emu(4400000))
tf_df = tb_df.text_frame
tf_df.word_wrap = True

pd_1 = tf_df.paragraphs[0]
pd_1.text = "🛡 BLINDAGEM TÉCNICA RDAA"
pd_1.font.name = FONT_LATO
pd_1.font.size = Pt(12)
pd_1.font.bold = True
pd_1.font.color.rgb = COLOR_ACCENT
pd_1.space_after = Pt(14)

pd_2 = tf_df.add_paragraph()
pd_2.text = "• Primazia do Art. 373, I: Incumbe com exclusividade à reconvinte comprovar aquisição na constância e com esforço comum.\n\n• Presunção Inoperante: O art. 1.660 CC não supre a inércia probatória primária de quem formula o pedido reconvencional.\n\n• Lastro Exclusivo de Reserva: Sílvio detém extratos bancários e contratos atestando aquisição autônoma fora da comunhão."
pd_2.font.name = FONT_LATO
pd_2.font.size = Pt(10)
pd_2.font.color.rgb = COLOR_DARK_SLATE

# ==========================================
# SLIDE 7: TIMELINE / ROADMAP
# ==========================================
s7 = prs.slides.add_slide(blank_layout)
add_header(s7, "Roadmap de Providências Imediatas", "Cronograma de Atuação")
add_footer(s7, 7, 7)

steps_7 = [
    ("ETAPA 1", "JULGAMENTO COLEGIADO TJMG", "Distribuir memoriais na 8ª Câmara Cível para confirmar a obrigação exclusiva do IPTU com a ocupante.", COLOR_ACCENT),
    ("ETAPA 2", "QUESITOS DE PERÍCIA LOCATÍCIA", "Apresentar quesitos técnicos na 10ª Cível para arbitrar o valor do aluguel da casa da Rua das Begônias.", COLOR_DARK_SLATE),
    ("ETAPA 3", "BAIXA DO PROTESTO", "Notificar a Fazenda Municipal e o Cartório de Protestos para certificar o cumprimento da liminar por Ana.", COLOR_DARK_SLATE),
    ("ETAPA 4", "EXCLUSÃO DO EDIFÍCIO SENSE", "Consolidar a incomunicabilidade na partilha da 3ª Vara de Família mediante prova bancária documental.", COLOR_DARK_SLATE)
]

w7 = Emu(2600000)
for idx, (tag, title, desc, col) in enumerate(steps_7):
    sx = Emu(594360 + idx * 2800000)
    box7 = s7.shapes.add_shape(MSO_SHAPE.RECTANGLE, sx, Emu(1200000), w7, Emu(4900000))
    box7.fill.solid()
    if idx == 0:
        box7.fill.fore_color.rgb = COLOR_ACCENT_LIGHT
        box7.line.color.rgb = COLOR_ACCENT
        box7.line.width = Pt(1.5)
    else:
        box7.fill.fore_color.rgb = COLOR_CARD_BG
        box7.line.color.rgb = COLOR_CARD_BORDER
        
    tb7 = s7.shapes.add_textbox(sx + Emu(160000), Emu(1400000), w7 - Emu(320000), Emu(4500000))
    tf7 = tb7.text_frame
    tf7.word_wrap = True
    
    pt = tf7.paragraphs[0]
    pt.text = tag
    pt.font.name = FONT_LATO
    pt.font.size = Pt(10)
    pt.font.bold = True
    pt.font.color.rgb = col
    
    ph = tf7.add_paragraph()
    ph.text = title
    ph.font.name = FONT_LATO
    ph.font.size = Pt(12)
    ph.font.bold = True
    ph.font.color.rgb = COLOR_DARK_SLATE
    ph.space_after = Pt(12)
    
    pd = tf7.add_paragraph()
    pd.text = desc
    pd.font.name = FONT_LATO
    pd.font.size = Pt(9.5)
    pd.font.color.rgb = COLOR_STRUCTURE

prs.save(OUTPUT_PPTX)
print(f"PPTX v3 (Lato + Visual Law + Bento + Diagramas) gerado em: {OUTPUT_PPTX}")
