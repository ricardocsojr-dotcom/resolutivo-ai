#!/usr/bin/env python3
"""
extrair_agenda_hoje.py — Extrator de Prazos Processuais para Resolutivo.AI Cockpit
================================================================================
Regras:
1. Filtro na coluna Adv: Somente o que começar com 'Ric' ou 'Flavia'/'Flávia' (e qualquer texto à frente).
2. Exclusão estrita de TAREFAS (ex: solicitar emissão de guia, subsídios, etc.). Mantém apenas PRAZOS judiciais.
3. Extrai Número do Processo E Tramitação Dona (TD).
4. Ordenação por Prazo Fatal crescente.
5. Exporta para agenda_hoje.json para consumo do Cockpit e do Plugin Hermes.
"""

import os
import glob
import json
import re
from datetime import datetime

AGENDA_DIR = r"C:\Users\ricar\OneDrive - RD\Resolutivo\001. AGENDA DIÁRIA"
OUTPUT_PLUGIN = r"C:\Users\ricar\AppData\Local\hermes\desktop-plugins\resolutivo-ai\agenda_hoje.json"
OUTPUT_REPO = r"C:\Projetos\resolutivo-ai\dados-juridicos\agenda_hoje.json"

TAREFAS_EXCLUIDAS = {
    "2 SOLICITAR EMISSÃO DE GUIA",
    "2 SOLICITAÇÃO DE SUBSÍDIO",
    "2 RETORNO DO SUBSÍDIO",
    "2 SOLICITAR CARTA DE PREPOSTO",
    "2 CONTRATAR PREPOSTO/ADVOGADO CORRRESPONDENTE",
    "2 JUNTAR CP/SUB/AVISAR TESTEMUNHA",
    "2 PROVIDÊNCIAS AUDIÊNCIA",
    "2 PAGAMENTO DO DÉBITO",
    "2 CUMPRIMENTO DA LIMINAR"
}

def get_latest_agenda():
    files = glob.glob(os.path.join(AGENDA_DIR, "Agenda *.xlsx"))
    if not files:
        return None
    
    def sort_key(f):
        m = re.search(r"Agenda\s+(\d{2})\.(\d{2})\.(\d{4})\.xlsx", os.path.basename(f), re.I)
        if m:
            return f"{m.group(3)}-{m.group(2)}-{m.group(1)}"
        return os.path.getmtime(f)
    
    return sorted(files, key=sort_key)[-1]

def extrair():
    latest_file = get_latest_agenda()
    if not latest_file:
        print("[ERRO] Nenhuma agenda encontrada.")
        return False

    import openpyxl
    wb = openpyxl.load_workbook(latest_file, data_only=True)
    filename = os.path.basename(latest_file)

    if "CPJ3C" not in wb.sheetnames:
        print("[ERRO] Aba CPJ3C não encontrada.")
        return False

    ws = wb["CPJ3C"]
    headers = [str(cell.value or "").strip() for cell in ws[1]]

    try:
        adv_idx = headers.index("Adv")
        prazo_idx = headers.index("Prazo")
        td_idx = headers.index("Tramitação Dona")
        fatal_idx = headers.index("Prazo Fatal")
        cliente_idx = headers.index("Cliente")
        proc_idx = headers.index("Número do processo")
        resp_idx = headers.index("Responsável do Processo") if "Responsável do Processo" in headers else adv_idx
    except ValueError as e:
        print("[ERRO] Coluna obrigatória não encontrada:", e)
        return False

    prazos = []

    for r in range(2, ws.max_row + 1):
        adv = str(ws.cell(r, adv_idx + 1).value or "").strip()
        prazo_raw = str(ws.cell(r, prazo_idx + 1).value or "").strip()

        # 1. Filtro Adv: Ric... ou Flavia... / Flávia...
        adv_lower = adv.lower()
        match_adv = adv_lower.startswith("ric") or adv_lower.startswith("flavia") or adv_lower.startswith("flávia")
        if not match_adv:
            continue

        # 2. Filtro Tarefas: exclui tarefas (guia, subsídio, etc.)
        prazo_upper = prazo_raw.upper()
        if prazo_raw in TAREFAS_EXCLUIDAS or "GUIA" in prazo_upper or "SUBSÍDIO" in prazo_upper or "SUBSIDIO" in prazo_upper:
            continue

        nome_prazo = re.sub(r"^2\s+", "", prazo_raw).strip()
        cliente = str(ws.cell(r, cliente_idx + 1).value or "").strip()
        processo = str(ws.cell(r, proc_idx + 1).value or "").strip()
        tramitacao = str(ws.cell(r, td_idx + 1).value or "").strip()
        fatal_raw = str(ws.cell(r, fatal_idx + 1).value or "").strip()
        responsavel = str(ws.cell(r, resp_idx + 1).value or "").strip()

        data_fatal_fmt = ""
        data_fatal_iso = ""
        if fatal_raw:
            data_fatal_iso = fatal_raw.split(" ")[0]
            try:
                dt = datetime.strptime(data_fatal_iso, "%Y-%m-%d")
                data_fatal_fmt = dt.strftime("%d/%m/%Y")
            except Exception:
                data_fatal_fmt = data_fatal_iso

        prazos.append({
            "prazo": nome_prazo,
            "cliente": cliente,
            "processo": processo,
            "tramitacao": tramitacao,
            "prazo_fatal": data_fatal_fmt,
            "prazo_fatal_iso": data_fatal_iso,
            "adv": adv,
            "responsavel": responsavel
        })

    # Ordena por data fatal crescente
    prazos.sort(key=lambda x: x["prazo_fatal_iso"] or "9999-99-99")

    data = {
        "arquivo": filename,
        "total_prazos": len(prazos),
        "data_extracao": datetime.now().strftime("%d/%m/%Y %H:%M"),
        "prazos": prazos
    }

    os.makedirs(os.path.dirname(OUTPUT_PLUGIN), exist_ok=True)
    with open(OUTPUT_PLUGIN, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    os.makedirs(os.path.dirname(OUTPUT_REPO), exist_ok=True)
    with open(OUTPUT_REPO, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)

    print(f"[OK] {len(prazos)} prazos judiciais com Tramitação extraídos de {filename}.")
    return True

if __name__ == "__main__":
    extrair()
