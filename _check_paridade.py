from pptx import Presentation

html = open(r'C:/Projetos/resolutivo-ai/rescisoria_trivale_minare_2slides.html', encoding='utf-8').read()
prs = Presentation(r'C:/Projetos/resolutivo-ai/rescisoria_trivale_minare_2slides.pptx')
pptx_texts = set()
for slide in prs.slides:
    for shp in slide.shapes:
        if shp.has_text_frame:
            t = shp.text_frame.text.strip()
            if t:
                pptx_texts.add(t)

checks = [
    'A quitação só está provada quando as 4 etapas se conectam.',
    '85% do valor tratado como dívida já estava pago e documentado.',
    'R$ 128.370,50',
    'R$ 18.523,21',
    'R$ 150.730,93',
]
for c in checks:
    in_html = c in html
    in_pptx = any(c in t for t in pptx_texts)
    print(f'{c[:55]!r:60} HTML={in_html}  PPTX={in_pptx}')
